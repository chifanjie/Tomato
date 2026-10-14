#!/usr/bin/env python3
"""Build the graph-layout PRD DOCX from the maintained Markdown source."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17202A"
MUTED = "667085"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D0D5DD"
WHITE = "FFFFFF"
GREEN = "168A43"
LIGHT_GREEN = "EAF8EF"
RED = "B42318"
LIGHT_RED = "FEEDEC"
BODY_FONT = "Hiragino Sans GB"
EAST_ASIA_FONT = "Hiragino Sans GB"
MONO_FONT = "Menlo"


def set_east_asia_font(run, font_name: str = EAST_ASIA_FONT) -> None:
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clear_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    r_fonts.set(qn("w:cs"), EAST_ASIA_FONT)
    r_pr.append(r_fonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def add_inline_runs(paragraph, text: str, *, font_size: float | None = None, color: str | None = None) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            run.font.name = BODY_FONT
            if font_size:
                run.font.size = Pt(font_size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            set_east_asia_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = BODY_FONT
            if font_size:
                run.font.size = Pt(font_size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            set_east_asia_font(run)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(font_size or 9.5)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            set_east_asia_font(run, MONO_FONT)
        else:
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        run.font.name = BODY_FONT
        if font_size:
            run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        set_east_asia_font(run)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal._element.rPr.rFonts.set(qn(f"w:{key}"), EAST_ASIA_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            style._element.rPr.rFonts.set(qn(f"w:{key}"), EAST_ASIA_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = BODY_FONT
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        caption._element.rPr.rFonts.set(qn(f"w:{key}"), EAST_ASIA_FONT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = BODY_FONT
        style.font.size = Pt(10.5)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            style._element.rPr.rFonts.set(qn(f"w:{key}"), EAST_ASIA_FONT)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.1


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.38)
        section.footer_distance = Inches(0.38)

        header = section.header
        table = header.add_table(rows=1, cols=2, width=Inches(6.5))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(4.6)
        table.columns[1].width = Inches(1.9)
        clear_table_borders(table)
        left = table.cell(0, 0).paragraphs[0]
        left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = left.add_run("UNDERSTAND CODE · PRODUCT REQUIREMENTS")
        run.font.name = BODY_FONT
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        set_east_asia_font(run)
        right = table.cell(0, 1).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = right.add_run("DRAFT · 2026-07-22")
        run.font.name = BODY_FONT
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        set_east_asia_font(run)

        footer = section.footer
        table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(5.2)
        table.columns[1].width = Inches(1.3)
        clear_table_borders(table)
        left = table.cell(0, 0).paragraphs[0]
        run = left.add_run("Graph Layout Quality Upgrade · Internal Review")
        run.font.name = BODY_FONT
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        set_east_asia_font(run)
        right = table.cell(0, 1).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = right.add_run("Page ")
        run.font.name = BODY_FONT
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        add_page_number(right)


def add_accent_bar(doc: Document, color: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    cell.height = Inches(0.06)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    clear_table_borders(table)


def add_callout(doc: Document, text: str, *, fill: str = LIGHT_BLUE, accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.10)
    table.columns[1].width = Inches(6.20)
    clear_table_borders(table)
    marker = table.cell(0, 0)
    set_cell_shading(marker, accent)
    set_cell_margins(marker, top=0, start=0, bottom=0, end=0)
    marker.paragraphs[0].paragraph_format.space_after = Pt(0)
    cell = table.cell(0, 1)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, text, font_size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_metric_strip(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [("846", "nodes"), ("985", "edges"), ("0", "calls in artifact"), ("3", "delivery phases")]
    for idx, (value, label) in enumerate(labels):
        cell = table.cell(0, idx)
        cell.width = Inches(1.625)
        set_cell_shading(cell, LIGHT_GRAY if idx % 2 == 0 else WHITE)
        set_cell_margins(cell, top=140, start=120, bottom=140, end=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        r.font.name = BODY_FONT
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(BLUE)
        set_east_asia_font(r)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(label)
        r.font.name = BODY_FONT
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        set_east_asia_font(r)
    clear_table_borders(table)


def add_cover(doc: Document) -> None:
    add_accent_bar(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    r.font.name = BODY_FONT
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    r.font.letter_spacing = Pt(1)
    set_east_asia_font(r)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("大规模代码图谱\n自动布局质量升级")
    r.font.name = BODY_FONT
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)
    set_east_asia_font(r)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Understand Code · 关系图谱 · v1.0")
    r.font.name = BODY_FONT
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    set_east_asia_font(r)

    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    values = [
        ("STATUS", "Draft for Review", "DATE", "2026-07-22"),
        ("OWNER", "Product / Frontend / Graph Analysis", "SCOPE", "Scanner + Dashboard + Layout"),
    ]
    for row_idx, row_values in enumerate(values):
        for col_idx, value in enumerate(row_values):
            cell = meta.cell(row_idx, col_idx)
            cell.width = Inches(1.625)
            set_cell_margins(cell, top=90, start=100, bottom=90, end=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            r.font.name = BODY_FONT
            r.font.size = Pt(8.5 if col_idx % 2 == 0 else 9.5)
            r.font.bold = col_idx % 2 == 0
            r.font.color.rgb = RGBColor.from_string(MUTED if col_idx % 2 == 0 else INK)
            set_east_asia_font(r)
            if col_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
    meta.style = "Table Grid"

    doc.add_paragraph()
    add_callout(
        doc,
        "决策摘要：P0 先修复展示集合与现有 Dagre；P1 增加可选且版本化的 topology.v1；P2 仅在同一批基准数据上显著优于 Dagre 且性能达标后，用 ELK Layered 替换布局核心。React Flow、Worker、缓存、下钻与 fallback 保持不变。",
    )
    add_metric_strip(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Layout comparison  ")
    r.bold = True
    r.font.name = BODY_FONT
    r.font.size = Pt(9.5)
    set_east_asia_font(r)
    p.add_run("当前布局 · Dagre 快修 · ELK 目标方案")


def add_contents(doc: Document, headings: list[str]) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.add_run("目录").bold = True
    p = doc.add_paragraph()
    p.add_run("评审时建议先阅读 2、6、7、8、9、11、12、14、15 节。").italic = True
    table = doc.add_table(rows=(len(headings) + 1) // 2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.15)
    table.columns[1].width = Inches(3.15)
    clear_table_borders(table)
    midpoint = (len(headings) + 1) // 2
    for idx, heading in enumerate(headings):
        col = 0 if idx < midpoint else 1
        row = idx if idx < midpoint else idx - midpoint
        cell = table.cell(row, col)
        set_cell_margins(cell, top=75, start=70, bottom=75, end=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(heading)
        r.font.name = BODY_FONT
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        set_east_asia_font(r)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        rows.append(cells)
        idx += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows, idx


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = rows[0]
    weights = [max(6, min(28, len(re.sub(r"[`*]", "", header[i] if i < len(header) else "")))) for i in range(cols)]
    total = sum(weights) or cols
    widths = [6.3 * w / total for w in weights]
    for col_idx, width in enumerate(widths):
        width = max(0.72, width)
        table.columns[col_idx].width = Inches(width)
    for row_idx, row in enumerate(rows):
        prevent_row_split(table.rows[row_idx])
        if row_idx == 0:
            set_repeat_table_header(table.rows[row_idx])
        for col_idx in range(cols):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            value = row[col_idx] if col_idx < len(row) else ""
            add_inline_runs(p, value, font_size=8.3, color=INK)
            for run in p.runs:
                run.bold = row_idx == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code.rstrip())
    run.font.name = MONO_FONT
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(INK)
    set_east_asia_font(run, MONO_FONT)
    clear_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(doc: Document, image_path: Path, alt: str) -> None:
    if not image_path.exists():
        p = doc.add_paragraph()
        r = p.add_run(f"[Missing image: {alt}]")
        r.font.color.rgb = RGBColor.from_string(RED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Inches(6.3))
    # python-docx does not expose image alt text through its public API.
    # Populate the underlying DrawingML metadata so screen readers receive
    # the same descriptive label used by the Markdown source.
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt)


PAGE_BREAK_H2 = ("5.", "6.", "7.", "8.", "9.", "11.", "12.", "16.", "附录")


def add_body(doc: Document, lines: list[str], source_dir: Path) -> None:
    idx = 0
    in_code = False
    code_lines: list[str] = []
    while idx < len(lines):
        raw = lines[idx].rstrip("\n")
        stripped = raw.strip()
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, "\n".join(code_lines))
                in_code = False
            idx += 1
            continue
        if in_code:
            code_lines.append(raw)
            idx += 1
            continue
        if not stripped:
            idx += 1
            continue
        if stripped.startswith("|") and idx + 1 < len(lines) and lines[idx + 1].strip().startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_markdown_table(doc, rows)
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = Path(image_match.group(2))
            if not image_path.is_absolute():
                image_path = source_dir / image_path
            add_image(doc, image_path.resolve(), image_match.group(1))
            idx += 1
            continue
        if stripped.startswith("_") and stripped.endswith("_"):
            p = doc.add_paragraph(style="Caption")
            add_inline_runs(p, stripped[1:-1], font_size=9, color=MUTED)
            idx += 1
            continue
        if stripped.startswith("> "):
            add_callout(doc, stripped[2:])
            idx += 1
            continue
        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1)) - 1
            text = heading_match.group(2)
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            if level == 1 and text.startswith(PAGE_BREAK_H2):
                p.paragraph_format.page_break_before = True
            add_inline_runs(p, text, font_size={1: 16, 2: 13, 3: 12}.get(level, 11), color=BLUE if level < 3 else DARK_BLUE)
            idx += 1
            continue
        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, bullet_match.group(1), font_size=10.5, color=INK)
            idx += 1
            continue
        number_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number_match:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, number_match.group(1), font_size=10.5, color=INK)
            idx += 1
            continue
        p = doc.add_paragraph()
        add_inline_runs(p, stripped.rstrip("  "), font_size=10.5, color=INK)
        next_idx = idx + 1
        while next_idx < len(lines) and not lines[next_idx].strip():
            next_idx += 1
        if next_idx < len(lines) and lines[next_idx].strip().startswith("```"):
            p.paragraph_format.keep_with_next = True
        idx += 1


def build(source: Path, output: Path) -> None:
    markdown = source.read_text(encoding="utf-8")
    all_lines = markdown.splitlines()
    body_start = next(i for i, line in enumerate(all_lines) if line.startswith("## 1. "))
    body_lines = all_lines[body_start:]
    h2_headings = [line[3:].strip() for line in body_lines if line.startswith("## ")]

    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    props = doc.core_properties
    props.title = "大规模代码图谱自动布局质量升级 PRD"
    props.subject = "Understand Code graph layout quality upgrade"
    props.author = "OpenAI Codex"
    props.keywords = "Understand Code, graph layout, Dagre, ELK, topology, PRD"
    props.comments = "Draft for internal review; generated from maintained Markdown source."

    add_cover(doc)
    doc.add_page_break()
    add_contents(doc, h2_headings)
    doc.add_page_break()
    add_body(doc, body_lines, source.parent)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_graph_layout_prd.py SOURCE.md OUTPUT.docx")
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
